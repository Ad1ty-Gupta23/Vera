"""
Turns an uploaded file's raw bytes into plain text.

Supported types: .pdf, .txt, .md, .docx (pypdf + python-docx are both
pure-Python and cheap to depend on, so DOCX is included).
"""
from __future__ import annotations

import io
import logging

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {"pdf", "txt", "md", "docx"}


class ExtractionError(ValueError):
    """Raised when a file can't be turned into usable text."""


def extension_for(filename: str) -> str:
    if "." not in filename:
        raise ExtractionError("File has no extension — can't tell what type it is.")
    return filename.rsplit(".", 1)[-1].lower()


def extract_text(filename: str, content: bytes) -> str:
    """
    Dispatches to the right extractor based on file extension.
    Raises ExtractionError for unsupported types or unreadable content —
    callers should catch this and mark the document "failed" rather than
    letting it bubble up as a 500.
    """
    ext = extension_for(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(
            f"Unsupported file type '.{ext}'. Supported types: "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )

    if ext in ("txt", "md"):
        text = _extract_plain_text(content)
    elif ext == "pdf":
        text = _extract_pdf(content)
    elif ext == "docx":
        text = _extract_docx(content)
    else:  # pragma: no cover — guarded by SUPPORTED_EXTENSIONS above
        raise ExtractionError(f"Unsupported file type '.{ext}'.")

    text = text.strip()
    if not text:
        raise ExtractionError(
            "No readable text found in this file — it may be empty, scanned "
            "images, or corrupted."
        )
    return text


def _extract_plain_text(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        # Best-effort fallback for files saved with a different encoding
        # rather than rejecting outright.
        return content.decode("latin-1", errors="replace")


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("PDF support isn't installed on the server.") from exc

    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise ExtractionError(f"Couldn't open this PDF: {exc}") from exc

    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001 — one bad page shouldn't fail the whole doc
            logger.warning("[knowledge] failed to extract a PDF page: %s", exc)
    return "\n\n".join(pages)


def _extract_docx(content: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("DOCX support isn't installed on the server.") from exc

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise ExtractionError(f"Couldn't open this DOCX file: {exc}") from exc

    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    # Tables (e.g. pricing tables, FAQ tables) carry real business info too.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return "\n".join(paragraphs)
